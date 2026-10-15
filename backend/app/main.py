from fastapi import FastAPI, UploadFile, File, HTTPException, Depends
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from datetime import datetime
from uuid import UUID
import os

from .config import settings
from .database import get_db, init_db
from .models import (
    Company, PaymentType, PaymentCategory, Payee,
    Document, ExtractedData, Requisition, LineItem,
)
from .services.mistral_service import EXPENSE_CATEGORIES
from .services.word_generator import WordGenerator
from .workers.processor import process_document


app = FastAPI(title=settings.APP_NAME)

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "https://reimbursement-frontend.onrender.com",
        "https://reimbursement-frontend-*.onrender.com",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    init_db()
    print(f"✅ Application started: {settings.APP_NAME}")
    print(f"🌐 API: http://localhost:8000")


# ============ SYSTEM ============

@app.get("/")
async def root():
    return {"name": settings.APP_NAME, "status": "running"}


@app.get("/health")
async def health():
    return {"status": "healthy"}


# ============ LOOKUPS ============

@app.get("/api/categories")
async def get_expense_categories():
    return {"categories": [c["name"] for c in EXPENSE_CATEGORIES]}


@app.get("/api/companies")
async def list_companies(db: Session = Depends(get_db)):
    rows = db.query(Company).filter(Company.is_active == True).all()
    return [
        {"id": str(c.id), "name": c.name, "short_name": c.short_name}
        for c in rows
    ]


@app.get("/api/companies/{company_id}/categories")
async def list_company_categories(
    company_id: UUID, db: Session = Depends(get_db)
):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    rows = (
        db.query(PaymentCategory)
        .filter(
            PaymentCategory.company_id == company_id,
            PaymentCategory.is_active == True,
        )
        .all()
    )

    grouped: dict = {}
    for c in rows:
        ptype = c.payment_type.name
        grouped.setdefault(ptype, []).append({
            "id": str(c.id),
            "name": c.name,
            "frequency": c.frequency,
        })

    order = ["Monthly Payment", "Annual Payment", "AD HOC"]
    return {k: grouped.get(k, []) for k in order if k in grouped}


@app.post("/api/companies/{company_id}/categories")
async def create_company_category(
    company_id: UUID,
    data: dict,
    db: Session = Depends(get_db),
):
    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    name = (data.get("name") or "").strip()
    payment_type_name = (data.get("payment_type") or "").strip()

    if not name:
        raise HTTPException(400, "name is required")
    if not payment_type_name:
        raise HTTPException(400, "payment_type is required")

    payment_type = (
        db.query(PaymentType)
        .filter(PaymentType.name == payment_type_name)
        .first()
    )
    if not payment_type:
        raise HTTPException(400, f"Invalid payment_type: {payment_type_name}")

    existing = (
        db.query(PaymentCategory)
        .filter(
            PaymentCategory.company_id == company.id,
            PaymentCategory.payment_type_id == payment_type.id,
            PaymentCategory.name == name,
        )
        .first()
    )
    if existing:
        return {
            "id": str(existing.id),
            "name": existing.name,
            "frequency": existing.frequency,
            "payment_type": payment_type.name,
        }

    frequency_map = {
        "Monthly Payment": "monthly",
        "Annual Payment": "annual",
        "AD HOC": "adhoc",
    }
    frequency = frequency_map.get(payment_type.name, "monthly")

    new_cat = PaymentCategory(
        company_id=company.id,
        payment_type_id=payment_type.id,
        name=name,
        frequency=frequency,
        is_active=True,
    )
    db.add(new_cat)
    db.commit()
    db.refresh(new_cat)

    print(f"➕ New Nature: {company.name} → {payment_type.name} → {name}")

    return {
        "id": str(new_cat.id),
        "name": new_cat.name,
        "frequency": new_cat.frequency,
        "payment_type": payment_type.name,
    }


@app.get("/api/payees")
async def list_payees(db: Session = Depends(get_db)):
    rows = (
        db.query(Payee)
        .filter(Payee.is_active == True)
        .order_by(Payee.name)
        .all()
    )
    return [
        {"id": str(p.id), "name": p.name, "type": p.payee_type}
        for p in rows
    ]


@app.post("/api/payees")
async def create_payee(data: dict, db: Session = Depends(get_db)):
    name = (data.get("name") or "").strip()
    if not name:
        raise HTTPException(400, "name is required")

    existing = db.query(Payee).filter(Payee.name == name).first()
    if existing:
        return {
            "id": str(existing.id),
            "name": existing.name,
            "type": existing.payee_type,
        }

    payee = Payee(name=name, payee_type=data.get("type", "other"))
    db.add(payee)
    db.commit()
    return {"id": str(payee.id), "name": payee.name, "type": payee.payee_type}


# ============ DOCUMENTS ============

@app.post("/api/documents/upload")
async def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    allowed = [
        "image/jpeg", "image/png", "image/tiff",
        "image/bmp", "image/webp", "application/pdf",
    ]
    if file.content_type not in allowed:
        raise HTTPException(400, f"File type not allowed. Allowed: {allowed}")

    content = await file.read()
    if not content:
        raise HTTPException(400, "Empty file")

    document = Document(
        filename=file.filename,
        file_path=None,
        file_data=content,
        file_size=len(content),
        mime_type=file.content_type,
        status="pending",
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    print(f"📄 Uploaded: {document.filename} ({len(content)} bytes)")

    try:
        process_document.delay(str(document.id))
        print(f"✅ Task queued for {document.id}")
    except Exception as e:
        print(f"⚠️ Celery enqueue failed: {e}")
        document.status = "failed"
        document.error_message = f"Queue error: {e}"
        db.commit()

    return {
        "document_id": str(document.id),
        "filename": file.filename,
        "status": "processing",
    }


@app.post("/api/documents/manual")
async def create_manual_document(db: Session = Depends(get_db)):
    """
    Create a blank Document so the Review page can be opened manually
    without uploading a file.
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    document = Document(
        filename=f"manual_{timestamp}",
        file_path=None,
        file_data=None,
        file_size=0,
        mime_type="application/x-manual",
        status="completed",
    )
    db.add(document)
    db.commit()
    db.refresh(document)

    extracted = ExtractedData(
        document_id=document.id,
        raw_text="",
        structured_data={
            "items": [],
            "payee_name": None,
            "document_type": "manual",
            "total_amount": None,
        },
        confidence_score=1.0,
        is_verified=False,
    )
    db.add(extracted)
    db.commit()

    print(f"✍️  Manual document created: {document.id}")

    return {
        "document_id": str(document.id),
        "status": "ready",
    }


@app.get("/api/documents")
async def list_documents(
    skip: int = 0,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    documents = (
        db.query(Document)
        .order_by(Document.upload_date.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )
    return [
        {
            "id": str(d.id),
            "filename": d.filename,
            "status": d.status,
            "error_message": d.error_message,
            "upload_date": d.upload_date.isoformat() if d.upload_date else None,
            "file_size": d.file_size,
            "mime_type": d.mime_type,
        }
        for d in documents
    ]


@app.get("/api/documents/{doc_id}/status")
async def get_status(doc_id: UUID, db: Session = Depends(get_db)):
    document = db.query(Document).filter(Document.id == doc_id).first()
    if not document:
        raise HTTPException(404, "Document not found")
    return {
        "document_id": str(document.id),
        "filename": document.filename,
        "status": document.status,
        "error_message": document.error_message,
    }


@app.get("/api/documents/{doc_id}/extracted")
async def get_extracted(doc_id: UUID, db: Session = Depends(get_db)):
    extracted = (
        db.query(ExtractedData)
        .filter(ExtractedData.document_id == doc_id)
        .first()
    )
    if not extracted:
        return {
            "status": "pending",
            "items": [],
            "payee_name": None,
            "document_type": None,
            "total_amount": None,
            "raw_text": "",
        }

    structured = extracted.structured_data or {}
    return {
        "items": structured.get("items", []),
        "payee_name": structured.get("payee_name"),
        "document_type": structured.get("document_type"),
        "total_amount": structured.get("total_amount"),
        "raw_text": extracted.raw_text or "",
        "confidence": extracted.confidence_score,
        "verified": extracted.is_verified,
    }


# ============ SUGGESTIONS (AUTO-FILL) ============

@app.get("/api/documents/{doc_id}/suggestions")
async def get_suggestions(doc_id: UUID, db: Session = Depends(get_db)):
    extracted = (
        db.query(ExtractedData)
        .filter(ExtractedData.document_id == doc_id)
        .first()
    )
    if not extracted:
        return _empty_suggestions()

    structured = extracted.structured_data or {}
    items = structured.get("items", [])
    ai_payee = (structured.get("payee_name") or "").strip() or None
    raw_text = (extracted.raw_text or "").lower()

    nature_name = _detect_nature_from_text(raw_text)
    canonical_payee = _resolve_payee(ai_payee)

    if not nature_name and canonical_payee:
        nature_name = _nature_for_payee(canonical_payee)

    if not nature_name:
        nature_name = "Reimbursement"

    matches = _find_nature_in_tree(nature_name)

    company_short = None
    payment_type_name = None
    final_payee = None

    preferred = None
    if canonical_payee:
        for m in matches:
            if m["payees"] and canonical_payee in m["payees"]:
                preferred = m
                break

    chosen = preferred or (matches[0] if matches else None)

    if chosen:
        company_short = chosen["company_short"]
        payment_type_name = chosen["payment_type"]
        payees = chosen["payees"]

        if len(payees) == 1:
            final_payee = payees[0]
        elif len(payees) > 1:
            if canonical_payee and canonical_payee in payees:
                final_payee = canonical_payee
            else:
                final_payee = payees[0]
        else:
            final_payee = canonical_payee

    if not company_short:
        company_short = "LSI"
        payment_type_name = payment_type_name or "Monthly Payment"
        final_payee = canonical_payee

    company = (
        db.query(Company).filter(Company.short_name == company_short).first()
    )

    payment_category = None
    if company and nature_name and payment_type_name:
        payment_category = (
            db.query(PaymentCategory)
            .join(PaymentType)
            .filter(
                PaymentCategory.company_id == company.id,
                PaymentCategory.name == nature_name,
                PaymentType.name == payment_type_name,
            )
            .first()
        )

    period = _infer_period(items)

    return {
        "company_id": str(company.id) if company else None,
        "company_name": company.name if company else None,
        "payment_type": payment_type_name,
        "payment_category_id": str(payment_category.id) if payment_category else None,
        "payment_category_name": payment_category.name if payment_category else None,
        "payee_name": final_payee,
        "ai_payee_name": ai_payee,
        "period": period,
        "confidence": 1.0 if (final_payee and period) else 0.5,
        "detected_nature": nature_name,
    }


# ============ REVIEW ============

@app.post("/api/documents/{doc_id}/review")
async def save_review(
    doc_id: UUID,
    data: dict,
    db: Session = Depends(get_db),
):
    items = data.get("items", [])

    company_id = _to_uuid(data.get("company_id"))
    payment_category_id = _to_uuid(data.get("payment_category_id"))
    payee_name = (data.get("payee_name") or "").strip()
    period = _normalize_period(data.get("period"))

    if not payee_name:
        raise HTTPException(400, "payee_name is required")
    if not company_id or not payment_category_id:
        raise HTTPException(400, "company_id and payment_category_id are required")

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    category = (
        db.query(PaymentCategory)
        .filter(PaymentCategory.id == payment_category_id)
        .first()
    )
    if not category:
        raise HTTPException(404, "Payment category not found")

    payee = db.query(Payee).filter(Payee.name == payee_name).first()
    if not payee:
        payee = Payee(name=payee_name, payee_type="other")
        db.add(payee)
        db.commit()

    extracted = (
        db.query(ExtractedData)
        .filter(ExtractedData.document_id == doc_id)
        .first()
    )
    if extracted:
        existing_structured = extracted.structured_data or {}
        extracted.structured_data = {
            "items": items,
            "payee_name": payee_name,
            "company_id": str(company_id),
            "payment_category_id": str(payment_category_id),
            "period": period,
            "document_type": existing_structured.get("document_type"),
            "total_amount": existing_structured.get("total_amount"),
        }
        extracted.is_verified = True
        db.commit()

    requisition = (
        db.query(Requisition)
        .filter(
            Requisition.company_id == company.id,
            Requisition.payment_category_id == category.id,
            Requisition.payee_name == payee_name,
            Requisition.period == period,
        )
        .first()
    )

    if requisition is None:
        requisition = Requisition(
            pr_number=_generate_pr_number(db),
            company_id=company.id,
            payment_category_id=category.id,
            payee_id=payee.id,
            company_name=company.name,
            category_name=category.name,
            payment_type_name=category.payment_type.name,
            payee_name=payee_name,
            period=period,
            frequency=category.frequency,
            applicant="Juliana",
            nature="Payment Requisition",
            total_amount=0.0,
            status="draft",
        )
        db.add(requisition)
        db.commit()
        print(
            f"✨ Created {requisition.pr_number} "
            f"({category.name} / {payee_name} / {period or 'no period'})"
        )
    else:
        print(f"➕ Appending to {requisition.pr_number}")

    db.query(LineItem).filter(LineItem.source_document_id == doc_id).delete()
    db.flush()

    remaining = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == requisition.id)
        .order_by(LineItem.item_number)
        .all()
    )
    for i, li in enumerate(remaining, start=1):
        li.item_number = i

    real_items = [it for it in items if not it.get("is_total")]

    doc_total = None
    if extracted and extracted.structured_data:
        doc_total = extracted.structured_data.get("total_amount")

    if not real_items and doc_total is not None and float(doc_total) > 0:
        real_items = [{
            "description": _default_total_description(
                nature_name=category.name,
                doc_type=(extracted.structured_data or {}).get("document_type"),
            ),
            "amount_hkd": float(doc_total),
            "date": None,
            "category": None,
            "is_total": False,
        }]

    start_num = len(remaining) + 1
    for idx, item in enumerate(real_items):
        amount = float(item.get("amount_hkd", 0) or 0)
        line_item = LineItem(
            requisition_id=requisition.id,
            source_document_id=doc_id,
            item_number=start_num + idx,
            date=_parse_date(item.get("date")),
            description=item.get("description", ""),
            amount_hkd=amount,
            category=item.get("category") or None,
        )
        db.add(line_item)

    db.flush()

    all_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == requisition.id)
        .all()
    )
    requisition.total_amount = sum(li.amount_hkd for li in all_items)
    db.commit()

    try:
        _regenerate_word(db, requisition)
    except Exception as e:
        print(f"⚠️ Word regeneration failed: {e}")

    return {
        "message": "Review saved",
        "items_count": len(real_items),
        "requisition_id": str(requisition.id),
        "pr_number": requisition.pr_number,
    }


@app.post("/api/documents/{doc_id}/reprocess")
async def reprocess_document(doc_id: UUID, db: Session = Depends(get_db)):
    document = db.query(Document).filter(Document.id == doc_id).first()
    if not document:
        raise HTTPException(404, "Document not found")

    document.status = "pending"
    document.error_message = None
    db.commit()

    process_document.delay(str(document.id))
    return {"message": "Reprocessing started"}


# ============ DASHBOARD ============

@app.get("/api/dashboard/stats")
async def get_dashboard_stats(db: Session = Depends(get_db)):
    """
    Dashboard summary:
      - Last month: amount paid + count
      - This month: amount paid + count
      - This year: amount paid + count
      - Breakdown: Company → Category → { count, total }  (this year)
    """
    from dateutil.relativedelta import relativedelta

    today = datetime.utcnow().date()
    prev_month_date = today.replace(day=1) - relativedelta(months=1)
    prev_month_key = prev_month_date.strftime("%Y-%m")
    current_month_key = today.strftime("%Y-%m")
    current_year_key = today.strftime("%Y")

    last_rows = (
        db.query(Requisition)
        .filter(Requisition.period == prev_month_key)
        .all()
    )
    last_count = len(last_rows)
    last_total = sum(r.total_amount or 0 for r in last_rows)

    month_rows = (
        db.query(Requisition)
        .filter(Requisition.period == current_month_key)
        .all()
    )
    month_count = len(month_rows)
    month_total = sum(r.total_amount or 0 for r in month_rows)

    year_rows = (
        db.query(Requisition)
        .filter(Requisition.period.like(f"{current_year_key}%"))
        .all()
    )
    year_count = len(year_rows)
    year_total = sum(r.total_amount or 0 for r in year_rows)

    by_company = {}
    for r in year_rows:
        company = r.company_name or "Unknown"
        category = r.category_name or "Uncategorized"
        amount = r.total_amount or 0

        by_company.setdefault(company, {
            "total": 0.0,
            "count": 0,
            "categories": {},
            "company_id": str(r.company_id) if r.company_id else None,
        })
        by_company[company]["total"] += amount
        by_company[company]["count"] += 1

        by_company[company]["categories"].setdefault(category, {
            "count": 0,
            "total": 0.0,
        })
        by_company[company]["categories"][category]["count"] += 1
        by_company[company]["categories"][category]["total"] += amount

    ordered = {}
    for company in sorted(by_company.keys()):
        entry = by_company[company]
        entry["total"] = round(entry["total"], 2)

        sorted_cats = dict(
            sorted(
                entry["categories"].items(),
                key=lambda kv: kv[1]["total"],
                reverse=True,
            )
        )
        for cat in sorted_cats:
            sorted_cats[cat]["total"] = round(sorted_cats[cat]["total"], 2)

        ordered[company] = {
            "total": entry["total"],
            "count": entry["count"],
            "company_id": entry["company_id"],
            "categories": sorted_cats,
        }

    return {
        "last_month": {
            "period": prev_month_key,
            "label": prev_month_date.strftime("%B %Y"),
            "count": last_count,
            "total": round(last_total, 2),
        },
        "this_month": {
            "period": current_month_key,
            "label": today.strftime("%B %Y"),
            "count": month_count,
            "total": round(month_total, 2),
        },
        "this_year": {
            "period": current_year_key,
            "label": current_year_key,
            "count": year_count,
            "total": round(year_total, 2),
        },
        "by_company_this_year": ordered,
    }


# ============ ANNUAL REPORT EXCEL ============

def _build_annual_report_sheet(wb, company_name: str, year: str, requisitions: list):
    """
    Add one worksheet to `wb` showing Category × Month for the given company & year.
    """
    import openpyxl
    from openpyxl.styles import Font, Border, Side, Alignment

    sheet_name = (company_name[:28] + "..") if len(company_name) > 31 else company_name
    ws = wb.create_sheet(title=sheet_name)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    right = Alignment(horizontal="right")
    center = Alignment(horizontal="center")

    months = [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ]

    ws.merge_cells("A1:N1")
    ws["A1"] = f"{company_name} — Annual Report {year}"
    ws["A1"].font = Font(bold=True, size=14)

    header_row = 3
    ws.cell(row=header_row, column=1, value="Category").font = bold
    for i, m in enumerate(months):
        c = ws.cell(row=header_row, column=2 + i, value=m)
        c.font = bold
        c.alignment = center
    total_col = 2 + len(months)
    ws.cell(row=header_row, column=total_col, value="Total").font = bold
    ws.cell(row=header_row, column=total_col).alignment = center

    grid: dict = {}
    unknown_month: dict = {}

    for r in requisitions:
        category = r.category_name or "Uncategorized"
        amount = float(r.total_amount or 0)

        _, month = _split_period(r.period)
        if month == "—" or not month.isdigit():
            unknown_month[category] = unknown_month.get(category, 0) + amount
            continue

        m_idx = int(month) - 1
        if m_idx < 0 or m_idx > 11:
            unknown_month[category] = unknown_month.get(category, 0) + amount
            continue

        grid.setdefault(category, [0.0] * 12)
        grid[category][m_idx] += amount

    row = header_row + 1
    for category in sorted(grid.keys()):
        ws.cell(row=row, column=1, value=category)
        amounts = grid[category]
        for i, amt in enumerate(amounts):
            cell = ws.cell(row=row, column=2 + i, value=round(amt, 2))
            cell.number_format = "#,##0.00"
            cell.alignment = right
        total_cell = ws.cell(
            row=row,
            column=total_col,
            value=f"=SUM(B{row}:M{row})",
        )
        total_cell.font = bold
        total_cell.number_format = "#,##0.00"
        total_cell.alignment = right
        row += 1

    if unknown_month:
        row += 1
        ws.cell(row=row, column=1, value="Items with no month").font = Font(italic=True)
        row += 1
        for category, amt in sorted(unknown_month.items()):
            ws.cell(row=row, column=1, value=category)
            ws.cell(row=row, column=total_col, value=round(amt, 2)).number_format = "#,##0.00"
            row += 1

    total_row = row + 1
    ws.cell(row=total_row, column=1, value="Total").font = bold
    for i in range(2, 2 + len(months)):
        col_letter = openpyxl.utils.get_column_letter(i)
        c = ws.cell(
            row=total_row,
            column=i,
            value=f"=SUM({col_letter}{header_row + 1}:{col_letter}{row})",
        )
        c.font = bold
        c.number_format = "#,##0.00"
        c.alignment = right
    grand_total = ws.cell(
        row=total_row,
        column=total_col,
        value=f"=SUM(N{header_row + 1}:N{row})",
    )
    grand_total.font = bold
    grand_total.number_format = "#,##0.00"
    grand_total.alignment = right

    ws.column_dimensions["A"].width = 32
    for i in range(2, 2 + len(months) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 12

    for r in range(header_row, total_row + 1):
        for c in range(1, total_col + 1):
            ws.cell(row=r, column=c).border = border


@app.get("/api/companies/{company_id}/annual-report")
async def company_annual_report(
    company_id: UUID,
    year: str | None = None,
    db: Session = Depends(get_db),
):
    """
    Download an Excel workbook for a single company showing
    Category × Month for the given year (defaults to current year).
    """
    import openpyxl
    import tempfile

    company = db.query(Company).filter(Company.id == company_id).first()
    if not company:
        raise HTTPException(404, "Company not found")

    year = (year or datetime.utcnow().strftime("%Y")).strip()

    rows = (
        db.query(Requisition)
        .filter(
            Requisition.company_id == company.id,
            Requisition.period.like(f"{year}%"),
        )
        .order_by(Requisition.created_date.asc())
        .all()
    )

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _build_annual_report_sheet(wb, company.name, year, rows)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    safe_name = "".join(
        c if c.isalnum() or c in " -_" else "_"
        for c in company.name
    )[:60]

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{safe_name} - Annual Report {year}.xlsx",
    )


@app.get("/api/dashboard/annual-report")
async def dashboard_annual_report(
    year: str | None = None,
    db: Session = Depends(get_db),
):
    """
    Download one workbook with one sheet per company (Category × Month),
    plus a Summary sheet totalling across all companies.
    """
    import openpyxl
    import tempfile
    from openpyxl.styles import Font, Border, Side, Alignment

    year = (year or datetime.utcnow().strftime("%Y")).strip()

    companies = (
        db.query(Company)
        .filter(Company.is_active == True)
        .order_by(Company.name)
        .all()
    )

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = f"Summary {year}"

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    right = Alignment(horizontal="right")
    center = Alignment(horizontal="center")

    months = [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ]

    summary.merge_cells("A1:N1")
    summary["A1"] = f"All Companies — Annual Report {year}"
    summary["A1"].font = Font(bold=True, size=14)

    header_row = 3
    summary.cell(row=header_row, column=1, value="Company").font = bold
    for i, m in enumerate(months):
        c = summary.cell(row=header_row, column=2 + i, value=m)
        c.font = bold
        c.alignment = center
    total_col = 2 + len(months)
    summary.cell(row=header_row, column=total_col, value="Total").font = bold
    summary.cell(row=header_row, column=total_col).alignment = center

    row = header_row + 1
    for company in companies:
        rows = (
            db.query(Requisition)
            .filter(
                Requisition.company_id == company.id,
                Requisition.period.like(f"{year}%"),
            )
            .all()
        )

        if rows:
            _build_annual_report_sheet(wb, company.name, year, rows)

        amounts = [0.0] * 12
        for r in rows:
            _, month = _split_period(r.period)
            if not month.isdigit():
                continue
            m_idx = int(month) - 1
            if 0 <= m_idx < 12:
                amounts[m_idx] += float(r.total_amount or 0)

        summary.cell(row=row, column=1, value=company.name)
        for i, amt in enumerate(amounts):
            cell = summary.cell(row=row, column=2 + i, value=round(amt, 2))
            cell.number_format = "#,##0.00"
            cell.alignment = right
        total_cell = summary.cell(
            row=row, column=total_col, value=f"=SUM(B{row}:M{row})"
        )
        total_cell.font = bold
        total_cell.number_format = "#,##0.00"
        total_cell.alignment = right
        row += 1

    total_row = row
    summary.cell(row=total_row, column=1, value="Total").font = bold
    for i in range(2, 2 + len(months)):
        col = openpyxl.utils.get_column_letter(i)
        c = summary.cell(
            row=total_row,
            column=i,
            value=f"=SUM({col}{header_row + 1}:{col}{row - 1})",
        )
        c.font = bold
        c.number_format = "#,##0.00"
        c.alignment = right
    gt = summary.cell(
        row=total_row, column=total_col, value=f"=SUM(N{header_row + 1}:N{row - 1})"
    )
    gt.font = bold
    gt.number_format = "#,##0.00"
    gt.alignment = right

    summary.column_dimensions["A"].width = 36
    for i in range(2, 2 + len(months) + 1):
        summary.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 12

    for r in range(header_row, total_row + 1):
        for c in range(1, total_col + 1):
            summary.cell(row=r, column=c).border = border

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Annual Report {year}.xlsx",
    )


# ============ REQUISITIONS ============

@app.get("/api/requisitions")
async def list_requisitions(
    company_id: UUID | None = None,
    category_id: UUID | None = None,
    payee_name: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
):
    q = db.query(Requisition)
    if company_id:
        q = q.filter(Requisition.company_id == company_id)
    if category_id:
        q = q.filter(Requisition.payment_category_id == category_id)
    if payee_name:
        q = q.filter(Requisition.payee_name == payee_name)
    if status:
        q = q.filter(Requisition.status == status)

    rows = q.order_by(Requisition.created_date.desc()).all()

    result = []
    for r in rows:
        item_count = (
            db.query(LineItem)
            .filter(LineItem.requisition_id == r.id)
            .count()
        )
        result.append({
            "id": str(r.id),
            "pr_number": r.pr_number,
            "company_name": r.company_name,
            "category_name": r.category_name,
            "payment_type_name": r.payment_type_name,
            "payee_name": r.payee_name,
            "period": r.period,
            "total_amount": r.total_amount,
            "status": r.status,
            "item_count": item_count,
            "created_date": r.created_date.isoformat() if r.created_date else None,
        })
    return result


@app.get("/api/requisitions/folders")
async def list_folder_view(
    company_id: UUID | None = None,
    db: Session = Depends(get_db),
):
    """
    Return a nested tree for the folder view:

      Company → Year → Month → Category → Payee → [ requisitions ]
    """
    q = db.query(Requisition)
    if company_id:
        q = q.filter(Requisition.company_id == company_id)

    rows = q.order_by(Requisition.created_date.desc()).all()

    tree: dict = {}
    for r in rows:
        company = r.company_name or "Unknown Company"
        year, month = _split_period(r.period)
        category = r.category_name or "Uncategorized"
        payee = r.payee_name or "Unknown Payee"

        tree.setdefault(company, {})
        tree[company].setdefault(year, {})
        tree[company][year].setdefault(month, {})
        tree[company][year][month].setdefault(category, {})
        tree[company][year][month][category].setdefault(payee, []).append({
            "id": str(r.id),
            "pr_number": r.pr_number,
            "payee_name": r.payee_name,
            "payment_type_name": r.payment_type_name,
            "total_amount": r.total_amount,
            "status": r.status,
            "period": r.period,
            "created_date": r.created_date.isoformat() if r.created_date else None,
            "item_count": (
                db.query(LineItem)
                .filter(LineItem.requisition_id == r.id)
                .count()
            ),
        })

    sorted_tree = {}
    for company in sorted(tree.keys()):
        sorted_tree[company] = {}
        for year in sorted(tree[company].keys(), reverse=True):
            sorted_tree[company][year] = {}
            for month in sorted(tree[company][year].keys(), reverse=True):
                sorted_tree[company][year][month] = {}
                for category in sorted(tree[company][year][month].keys()):
                    sorted_tree[company][year][month][category] = {}
                    for payee in sorted(tree[company][year][month][category].keys()):
                        prs = sorted(
                            tree[company][year][month][category][payee],
                            key=lambda x: x.get("created_date") or "",
                            reverse=True,
                        )
                        sorted_tree[company][year][month][category][payee] = prs

    return sorted_tree


@app.get("/api/requisitions/{req_id}")
async def get_requisition(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Requisition not found")

    line_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == req_id)
        .order_by(LineItem.item_number)
        .all()
    )

    return {
        "id": str(req.id),
        "pr_number": req.pr_number,
        "company_name": req.company_name,
        "category_name": req.category_name,
        "payment_type_name": req.payment_type_name,
        "payee_name": req.payee_name,
        "applicant": req.applicant,
        "nature": req.nature,
        "period": req.period,
        "frequency": req.frequency,
        "total_amount": req.total_amount,
        "status": req.status,
        "created_date": req.created_date.isoformat() if req.created_date else None,
        "line_items": [
            {
                "id": str(li.id),
                "item_number": li.item_number,
                "date": li.date.isoformat() if li.date else None,
                "description": li.description,
                "amount_hkd": li.amount_hkd,
                "category": li.category,
            }
            for li in line_items
        ],
    }

@app.put("/api/requisitions/{req_id}")
async def update_requisition(
    req_id: UUID,
    data: dict,
    db: Session = Depends(get_db),
):
    """
    Fully update a requisition. Any of these fields may be included in the body:

      {
        "pr_number": "LSI-PR-2026-005",
        "company_id": "<uuid>",
        "payment_category_id": "<uuid>",
        "payee_name": "Gu Bai Chuan Cole",
        "period": "2026-07",
        "line_items": [
            {"date": "2026-07-15", "description": "...", "amount_hkd": 100.00, "category": "Hotel"},
            ...
        ]
      }

    Unspecified fields are left unchanged.
    """
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Requisition not found")

    # ---------------- PR number ----------------
    if "pr_number" in data:
        new_pr = (data.get("pr_number") or "").strip()
        if not new_pr:
            raise HTTPException(400, "pr_number cannot be empty")
        if new_pr != req.pr_number:
            clash = (
                db.query(Requisition)
                .filter(Requisition.pr_number == new_pr, Requisition.id != req.id)
                .first()
            )
            if clash:
                raise HTTPException(400, f"PR number '{new_pr}' already exists")
            req.pr_number = new_pr

    # ---------------- Company + Category ----------------
    new_company_id = _to_uuid(data.get("company_id"))
    new_category_id = _to_uuid(data.get("payment_category_id"))

    if new_company_id:
        new_company = db.query(Company).filter(Company.id == new_company_id).first()
        if not new_company:
            raise HTTPException(404, "Company not found")
        req.company_id = new_company.id
        req.company_name = new_company.name

    if new_category_id:
        new_category = (
            db.query(PaymentCategory)
            .filter(PaymentCategory.id == new_category_id)
            .first()
        )
        if not new_category:
            raise HTTPException(404, "Payment category not found")
        req.payment_category_id = new_category.id
        req.category_name = new_category.name
        req.payment_type_name = new_category.payment_type.name
        req.frequency = new_category.frequency

    # ---------------- Payee ----------------
    if "payee_name" in data:
        new_payee_name = (data.get("payee_name") or "").strip()
        if not new_payee_name:
            raise HTTPException(400, "payee_name cannot be empty")
        payee = db.query(Payee).filter(Payee.name == new_payee_name).first()
        if not payee:
            payee = Payee(name=new_payee_name, payee_type="other")
            db.add(payee)
            db.flush()
        req.payee_id = payee.id
        req.payee_name = new_payee_name

    # ---------------- Period ----------------
    if "period" in data:
        req.period = _normalize_period(data.get("period"))

    # ---------------- Line items ----------------
    if "line_items" in data and data["line_items"] is not None:
        line_items = data["line_items"]

        # Delete existing line items
        db.query(LineItem).filter(LineItem.requisition_id == req.id).delete()
        db.flush()

        # Add new ones
        for idx, item in enumerate(line_items):
            amount = float(item.get("amount_hkd", 0) or 0)
            line_item = LineItem(
                requisition_id=req.id,
                source_document_id=item.get("source_document_id") or None,
                item_number=idx + 1,
                date=_parse_date(item.get("date")),
                description=item.get("description", "") or "",
                amount_hkd=amount,
                category=item.get("category") or None,
            )
            db.add(line_item)

        db.flush()

        # Recompute total = sum of line items
        all_items = (
            db.query(LineItem)
            .filter(LineItem.requisition_id == req.id)
            .all()
        )
        req.total_amount = sum(li.amount_hkd for li in all_items)

    db.commit()
    db.refresh(req)

    # Regenerate the Word document with the new values
    try:
        _regenerate_word(db, req)
    except Exception as e:
        print(f"⚠️ Word regeneration failed: {e}")

    return {
        "message": "Requisition updated",
        "id": str(req.id),
        "pr_number": req.pr_number,
        "total_amount": req.total_amount,
    }

@app.delete("/api/requisitions/{req_id}")
async def delete_requisition(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Requisition not found")

    db.query(LineItem).filter(LineItem.requisition_id == req_id).delete()
    db.delete(req)
    db.commit()
    return {"message": "Deleted successfully"}


@app.get("/api/requisitions/{req_id}/export/word")
async def export_word(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Not found")

    docx_path = _regenerate_word(db, req)

    return FileResponse(
        docx_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{req.pr_number}.docx",
    )


@app.get("/api/requisitions/{req_id}/export/excel")
async def export_excel(req_id: UUID, db: Session = Depends(get_db)):
    req = db.query(Requisition).filter(Requisition.id == req_id).first()
    if not req:
        raise HTTPException(404, "Not found")

    line_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == req_id)
        .order_by(LineItem.item_number)
        .all()
    )

    import openpyxl
    from openpyxl.styles import Font, Border, Side, Alignment
    import tempfile

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{req.pr_number}_Detail List"

    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    period_label = req.period or ""
    payee_label = req.payee_name or ""
    title = f"Mr {payee_label}'s Reimbursement - {period_label}"
    ws.merge_cells("A1:K1")
    ws["A1"] = title
    ws["A1"].font = bold

    headers = ["", "Item", "No.", "Date", "Description", "RMB", "HKD", "RMB", "HKD", "HKD"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = bold

    groups: dict = {}
    for li in line_items:
        cat = (li.category or "Miscellaneous").strip() or "Miscellaneous"
        groups.setdefault(cat, []).append(li)

    row = 4
    for category, items in groups.items():
        cat_start_row = row
        ws.cell(row=row, column=2, value=category)

        for idx, li in enumerate(items):
            ws.cell(row=row, column=3, value=idx + 1)
            ws.cell(row=row, column=4, value=li.date if li.date else None)
            if li.date:
                ws.cell(row=row, column=4).number_format = "yyyy-mm-dd"
            desc_cell = ws.cell(row=row, column=5, value=li.description)
            desc_cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=7, value=li.amount_hkd)
            row += 1

        cat_end_row = row - 1
        ws.cell(
            row=cat_start_row,
            column=10,
            value=f"=SUM(G{cat_start_row}:G{cat_end_row})",
        )

    grand_total_row = row + 1
    ws.cell(row=grand_total_row, column=6, value=f"=SUM(G4:G{row-1})")
    ws.cell(row=grand_total_row, column=10, value=f"=SUM(J4:J{row-1})")

    subtotal_row = grand_total_row + 2
    ws.cell(row=subtotal_row, column=5, value="HKD sub-total")
    ws.cell(row=subtotal_row, column=7, value=f"=SUM(G4:G{row-1})")

    total_row = subtotal_row + 2
    ws.cell(row=total_row, column=5, value="Total")
    ws.cell(row=total_row, column=6, value=f"=F{subtotal_row}+G{subtotal_row}")

    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 6
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 60
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 12
    ws.column_dimensions["J"].width = 12
    ws.column_dimensions["K"].width = 20

    max_col = 10
    max_row = total_row
    for r in range(3, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            if cell.value is not None or (4 <= r <= row - 1):
                cell.border = border

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    wb.save(tmp.name)

    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"{req.pr_number}_Detail List.xlsx",
    )


# ============ HELPERS ============

def _empty_suggestions():
    return {
        "company_id": None, "payment_type": None,
        "payment_category_id": None, "payee_name": None, "period": None,
        "confidence": 0.0,
    }


def _normalize_name(s: str) -> str:
    import re as _re
    s = (s or "").lower()
    s = _re.sub(r"\b(mr|mrs|ms|dr|prof|professor)\.?\b", " ", s)
    s = _re.sub(r"[.,\-_/\\()\[\]]", " ", s)
    return " ".join(s.split())


def _resolve_payee(ai_payee: str | None) -> str | None:
    if not ai_payee:
        return None
    needle = _normalize_name(ai_payee)

    for alias, canonical in settings.PAYEE_ALIASES.items():
        if _normalize_name(alias) == needle:
            return canonical

    for alias, canonical in settings.PAYEE_ALIASES.items():
        a = _normalize_name(alias)
        if len(a) >= 4 and (a in needle or needle in a):
            return canonical

    return ai_payee.strip()


def _detect_nature_from_text(text: str) -> str | None:
    if not text:
        return None
    text_lower = text.lower()
    scores: dict = {}
    for nature, keywords in settings.NATURE_KEYWORDS.items():
        hits = sum(1 for kw in keywords if kw.lower() in text_lower)
        if hits > 0:
            scores[nature] = hits
    if not scores:
        return None
    return max(scores.items(), key=lambda kv: (kv[1], len(kv[0])))[0]


def _nature_for_payee(canonical_payee: str) -> str | None:
    preferred_order = settings.NATURE_PREFERRED_ORDER
    for nature in preferred_order:
        for comp_short, types in settings.COMPANY_TREE.items():
            for ptype, natures in types.items():
                if nature in natures and canonical_payee in natures[nature]:
                    return nature
    return None


def _find_nature_in_tree(nature_name: str) -> list:
    results = []
    for comp_short, types in settings.COMPANY_TREE.items():
        for ptype, natures in types.items():
            if nature_name in natures:
                results.append({
                    "company_short": comp_short,
                    "payment_type": ptype,
                    "payees": list(natures[nature_name]),
                })
    return results


def _regenerate_word(db: Session, req: Requisition) -> str:
    line_items = (
        db.query(LineItem)
        .filter(LineItem.requisition_id == req.id)
        .order_by(LineItem.item_number)
        .all()
    )

    company = None
    if req.company_id:
        company = db.query(Company).filter(Company.id == req.company_id).first()

    doc_data = {
        "requisition_id": str(req.id),
        "pr_number": req.pr_number,
        "applicant": req.applicant,
        "payee_name": req.payee_name,
        "nature": req.nature,
        "period": req.period,
        "category_name": req.category_name,
        "company_name": company.name if company else req.company_name,
        "company_address": company.address if company else None,
        "company_contact": None,
        "line_items": [
            {
                "date": li.date.strftime("%Y-%m-%d") if li.date else "",
                "description": li.description,
                "amount_hkd": li.amount_hkd,
                "category": li.category,
            }
            for li in line_items
        ],
    }
    return WordGenerator().generate_requisition_docx(doc_data)


def _generate_pr_number(db: Session) -> str:
    year = datetime.utcnow().year
    last = (
        db.query(Requisition)
        .order_by(Requisition.created_date.desc())
        .first()
    )
    next_num = 1
    if last:
        parts = last.pr_number.split("-")
        if len(parts) >= 4:
            try:
                next_num = int(parts[-1]) + 1
            except ValueError:
                next_num = 1
    return f"LSI-PR-{year}-{next_num:03d}"


def _to_uuid(value):
    if value is None or value == "":
        return None
    if isinstance(value, UUID):
        return value
    try:
        return UUID(str(value))
    except (ValueError, AttributeError, TypeError):
        return None


def _infer_period(items) -> str | None:
    valid = []
    current_year = datetime.utcnow().year
    for item in items:
        d = _parse_date(item.get("date"))
        if not d:
            continue
        if d.year < current_year - 2 or d.year > current_year + 1:
            continue
        valid.append(d)
    if valid:
        return min(valid).strftime("%Y-%m")
    return None


def _parse_date(date_str):
    if not date_str:
        return None
    if isinstance(date_str, datetime):
        return date_str

    for fmt in (
        "%Y-%m-%d", "%Y/%m/%d", "%d-%m-%Y", "%d/%m/%Y",
        "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%S.%f",
    ):
        try:
            return datetime.strptime(str(date_str).strip()[:19], fmt)
        except (ValueError, TypeError):
            continue
    return None


def _split_period(period: str | None) -> tuple:
    """
    Return (year, month) from a period string.
      "2026-07"       → ("2026", "07")
      "2026"          → ("2026", "—")
      "2026-07-15"    → ("2026", "07")
      None / garbage  → ("—", "—")
    """
    if not period or not isinstance(period, str):
        return ("—", "—")

    s = period.strip()

    if len(s) >= 7 and s[4] == "-":
        year = s[:4]
        month = s[5:7]
        if year.isdigit() and month.isdigit():
            return (year, month)

    if len(s) == 4 and s.isdigit():
        return (s, "—")

    return ("—", "—")


def _normalize_period(value):
    """
    Normalize a period string:
      "2026-07"     → "2026-07"
      "2026/07"     → "2026-07"
      "2026-7"      → "2026-07"
      "2026"        → "2026"
      "2026-07-15"  → "2026-07"
      "" / None     → None
    """
    if not value:
        return None

    s = str(value).strip().replace("/", "-").replace(".", "-")

    import re as _re

    m = _re.match(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?", s)
    if m:
        y, mo = m.group(1), int(m.group(2))
        return f"{y}-{mo:02d}"

    if _re.match(r"^\d{4}$", s):
        return s

    return s or None


def _default_total_description(nature_name: str, doc_type: str | None = None) -> str:
    """
    Build a sensible description when a document only has a total and no
    line items. Used to synthesize a single line item from the total.
    """
    nature_lower = (nature_name or "").lower()
    doc_lower = (doc_type or "").lower()

    if "reimbursement" in nature_lower:
        return "Reimbursement total"
    if "mpf" in nature_lower:
        return "MPF contribution (total)"
    if "salary" in nature_lower:
        return "Salary payment"
    if "group medical" in nature_lower:
        return "Group Medical insurance premium"
    if "employment compensation" in nature_lower:
        return "Employment Compensation insurance"
    if "br renewal" in nature_lower:
        return "Business Registration renewal fee"
    if "auditor" in nature_lower:
        return "Auditor services"
    if "legal" in nature_lower:
        return "Legal consultation fee"
    if "company secretary" in nature_lower:
        return "Company Secretary fee"
    if "electricity" in nature_lower:
        return "Electricity bill"
    if "management fee" in nature_lower:
        return "Management fee"
    if "cleaning" in nature_lower:
        return "Office cleaning service"
    if "credit card" in nature_lower:
        return "Credit card payment"
    if "microsoft" in nature_lower:
        return "Microsoft 365 subscription"
    if "software" in nature_lower:
        return "Software annual fee"
    if doc_lower == "cheque":
        return "Cheque payment"
    return "Payment"