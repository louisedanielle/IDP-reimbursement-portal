import openpyxl
from openpyxl.styles import Font, Alignment
from datetime import datetime
from pathlib import Path

from ..config import settings


class ExcelGenerator:
    def __init__(self):
        self.output_dir = Path(settings.OUTPUT_DIR)
        self.output_dir.mkdir(exist_ok=True)
    
    def generate_reimbursement_excel(self, data: dict) -> str:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"LSI-PR-{data['pr_number']}"
        
        # Title
        ws.merge_cells('A1:J2')
        ws['A1'] = "PAYMENT REQUISITION"
        ws['A1'].font = Font(size=16, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')
        
        # Info
        ws['A4'] = "PR No.:"
        ws['C4'] = data['pr_number']
        ws['A5'] = "Applicant:"
        ws['C5'] = data['applicant']
        ws['A6'] = "Nature:"
        ws['C6'] = "Reimbursement"
        ws['A8'] = "Name of Payee:"
        ws['C8'] = data['payee_name']
        ws['A9'] = "Amount payable:"
        ws['C9'] = f"HK${sum(i.get('amount_hkd', 0) for i in data['line_items']):.2f}"
        
        # Headers
        ws['A12'] = 'Date'
        ws['A12'].font = Font(bold=True)
        ws['C12'] = 'Particular'
        ws['C12'].font = Font(bold=True)
        ws['D12'] = 'Expenses in HK$'
        ws['D12'].font = Font(bold=True)
        
        # Items
        row = 14
        for item in data.get('line_items', []):
            ws[f'A{row}'] = item.get('date', '')
            ws[f'C{row}'] = item.get('description', '')
            ws[f'D{row}'] = item.get('amount_hkd', 0)
            row += 1
        
        # Total
        total_row = row + 1
        ws[f'D{total_row}'] = f'=SUM(D14:D{total_row-1})'
        ws[f'D{total_row}'].font = Font(bold=True)
        
        # Signature
        sign_row = total_row + 3
        ws[f'A{sign_row}'] = data['applicant']
        ws[f'C{sign_row}'] = 'Approved by:'
        
        # Column widths
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 20
        
        # Save
        filename = f"PR_{data['pr_number']}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = self.output_dir / filename
        wb.save(filepath)
        
        return str(filepath)